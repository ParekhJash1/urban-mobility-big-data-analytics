from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "Urban_Mobility_Big_Data_Analytics_Jash_Parekh.pptx"

NAVY = RGBColor(16, 43, 58)
INK = RGBColor(16, 32, 44)
TEAL = RGBColor(25, 124, 115)
MINT = RGBColor(138, 215, 207)
PALE_MINT = RGBColor(223, 241, 237)
BLUE = RGBColor(45, 117, 230)
ORANGE = RGBColor(240, 120, 68)
GREEN = RGBColor(24, 168, 120)
PURPLE = RGBColor(135, 101, 220)
WHITE = RGBColor(255, 255, 255)
PAPER = RGBColor(246, 249, 248)
MUTED = RGBColor(100, 115, 125)
LINE = RGBColor(220, 229, 232)
LIGHT_BLUE = RGBColor(232, 240, 253)
LIGHT_ORANGE = RGBColor(253, 239, 232)

W = Inches(13.333)
H = Inches(7.5)


def rgb(fill, transparency=0):
    fill.solid()
    fill.fore_color.rgb = fill
    if transparency:
        fill.transparency = transparency


def shape(slide, kind, x, y, w, h, fill=None, line=None, radius=False):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid()
        item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(0.8)
    return item


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         italic=False, margin=0.04, fit=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if fit:
        tf.word_wrap = True
    return box


def rich_text(slide, runs, x, y, w, h, size=14, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    for value, bold, run_color in runs:
        run = p.add_run()
        run.text = value
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = run_color or color
    return box


def line(slide, x1, y1, x2, y2, color=LINE, width=1.3, arrow=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if arrow:
        connector.line.end_arrowhead = True
    return connector


def base(slide, number, dark=False):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else PAPER
    if not dark:
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 0.12, TEAL)
    text(slide, "URBAN MOBILITY · BIG DATA ANALYTICS", 0.55, 7.15, 5.6, 0.18, 7.5, MINT if dark else MUTED, True, "Aptos")
    text(slide, f"{number:02d}", 12.2, 7.12, 0.55, 0.2, 8, MINT if dark else MUTED, True, "Aptos", PP_ALIGN.RIGHT)


def title(slide, kicker, heading, subtitle=None, dark=False):
    text(slide, kicker.upper(), 0.62, 0.52, 5.5, 0.2, 9, MINT if dark else TEAL, True, "Aptos")
    text(slide, heading, 0.6, 0.82, 11.8, 0.55, 28, WHITE if dark else INK, True, "Aptos Display")
    if subtitle:
        text(slide, subtitle, 0.62, 1.43, 11.2, 0.35, 11, RGBColor(181, 201, 205) if dark else MUTED)


def bullet_list(slide, items: Iterable[str], x, y, w, h, size=15, color=INK, gap=0.48):
    for i, item in enumerate(items):
        yy = y + i * gap
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, yy + 0.12, 0.10, 0.10, TEAL)
        text(slide, item, x + 0.22, yy, w - 0.22, 0.33, size, color)


def label_box(slide, label, x, y, w, fill, color=INK, size=12):
    item = shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.53, fill, fill)
    text(slide, label, x + 0.08, y + 0.04, w - 0.16, 0.42, size, color, True, "Aptos", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return item


def add_kpi(slide, label, value, note, x, y, accent):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 2.75, 1.33, WHITE, LINE)
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, 0.07, 1.33, accent)
    text(slide, label.upper(), x + 0.22, y + 0.18, 2.3, 0.18, 8, MUTED, True)
    text(slide, value, x + 0.22, y + 0.43, 2.3, 0.38, 22, INK, True, "Aptos Display")
    text(slide, note, x + 0.22, y + 0.96, 2.3, 0.18, 8, MUTED)


def add_bar_chart(slide, values, labels, x, y, w, h, color=BLUE, max_label=None):
    max_value = max(values) if values else 1
    chart_w = w - 0.55
    chart_h = h - 0.42
    line(slide, x + 0.42, y + 0.1, x + 0.42, y + h - 0.23, LINE, 1)
    line(slide, x + 0.42, y + h - 0.23, x + w, y + h - 0.23, LINE, 1)
    slot = chart_w / len(values)
    for i, value in enumerate(values):
        bh = chart_h * value / max_value
        bx = x + 0.53 + i * slot + slot * 0.17
        bw = slot * 0.62
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, bx, y + h - 0.23 - bh, bw, bh, color, color)
        text(slide, labels[i], bx - 0.04, y + h - 0.16, bw + 0.08, 0.2, 7, MUTED, False, "Aptos", PP_ALIGN.CENTER)
    if max_label:
        text(slide, max_label, x + 0.45, y + 0.02, 1.2, 0.18, 7, MUTED)


def add_line_chart(slide, values, labels, x, y, w, h, color=BLUE):
    left, bottom = x + 0.48, y + h - 0.28
    top, right = y + 0.12, x + w
    line(slide, left, top, left, bottom, LINE, 1)
    line(slide, left, bottom, right, bottom, LINE, 1)
    step = (right - left) / (len(values) - 1)
    max_value = max(values)
    points = []
    for i, value in enumerate(values):
        px = left + i * step
        py = bottom - (bottom - top - 0.25) * value / max_value
        points.append((px, py))
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        line(slide, x1, y1, x2, y2, color, 2.2)
    label_positions = {
        round(i * (len(values) - 1) / max(1, len(labels) - 1)): label
        for i, label in enumerate(labels)
    }
    for i, (px, py) in enumerate(points):
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, px - 0.045, py - 0.045, 0.09, 0.09, color, WHITE)
        if i in label_positions:
            text(slide, label_positions[i], px - 0.2, bottom + 0.05, 0.4, 0.18, 7, MUTED, False, "Aptos", PP_ALIGN.CENTER)


def add_code(slide, code, x, y, w, h, font_size=10):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, NAVY, NAVY)
    text(slide, code, x + 0.18, y + 0.12, w - 0.36, h - 0.22, font_size, MINT, False, "Courier New")


def build_presentation():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    base(slide, 1, True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 9.7, -1.4, 5.2, 5.2, NAVY, MINT)
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 10.35, -0.75, 3.9, 3.9, NAVY, RGBColor(46, 91, 105))
    text(slide, "BIG DATA ANALYTICS PROJECT", 0.75, 0.8, 5.8, 0.25, 11, MINT, True)
    text(slide, "Urban Mobility\nAnalytics Dashboard", 0.72, 1.35, 8.8, 1.45, 42, WHITE, True, "Aptos Display")
    text(slide, "An end-to-end data platform using Apache Spark, MongoDB, FastAPI, and a web dashboard", 0.78, 3.12, 7.1, 0.55, 15, RGBColor(181, 201, 205))
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.78, 4.45, 3.65, 1.14, TEAL, TEAL)
    text(slide, "JASH PAREKH", 1.03, 4.67, 3.1, 0.25, 16, WHITE, True)
    text(slide, "Roll No. 23BIT033", 1.03, 5.02, 3.1, 0.22, 11, MINT, True)
    text(slide, "January 2024 NYC TLC Yellow Taxi Records", 0.8, 6.42, 6.4, 0.2, 10, RGBColor(181, 201, 205))
    text(slide, "UNIVERSITY COURSE PRESENTATION", 9.0, 6.42, 3.5, 0.2, 8, MINT, True, "Aptos", PP_ALIGN.RIGHT)

    # 2. Executive summary
    slide = prs.slides.add_slide(blank)
    base(slide, 2)
    title(slide, "01 · project overview", "Turning millions of trips into decisions", "The project demonstrates a complete analytics lifecycle, from raw data to an interactive presentation layer.")
    add_kpi(slide, "Valid trips", "2.87M", "Spark-cleaned records", 0.65, 2.08, BLUE)
    add_kpi(slide, "Revenue", "$78.42M", "Aggregated total_amount", 3.62, 2.08, ORANGE)
    add_kpi(slide, "Average fare", "$18.49", "Per completed trip", 6.59, 2.08, GREEN)
    add_kpi(slide, "Avg. distance", "3.29 mi", "Trip distance", 9.56, 2.08, PURPLE)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.65, 3.83, 12.03, 2.27, PALE_MINT, PALE_MINT)
    text(slide, "CORE IDEA", 0.95, 4.14, 1.5, 0.2, 9, TEAL, True)
    text(slide, "Raw taxi records are too detailed to serve directly to a dashboard. Spark performs the expensive distributed work once, MongoDB stores compact analytical results, and FastAPI serves fast responses to the frontend.", 0.95, 4.48, 7.65, 0.85, 18, INK, True, "Aptos Display")
    text(slide, "Business questions answered", 9.05, 4.12, 2.8, 0.2, 9, TEAL, True)
    bullet_list(slide, ["When is demand highest?", "Which zones and routes lead?", "What is the revenue profile?"], 9.05, 4.48, 3.15, 1.25, 11, INK, 0.34)

    # 3. Problem and objectives
    slide = prs.slides.add_slide(blank)
    base(slide, 3)
    title(slide, "02 · problem definition", "Why this analytics system is needed", "Urban transport data is high-volume, time-sensitive, and difficult to analyze efficiently without a distributed workflow.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 2.0, 5.72, 4.32, WHITE, LINE)
    text(slide, "THE CHALLENGE", 1.0, 2.35, 2, 0.2, 9, ORANGE, True)
    bullet_list(slide, [
        "Millions of trip records contain invalid, missing, or unrealistic values.",
        "Repeatedly querying raw files creates unnecessary processing cost.",
        "Decision-makers need simple answers by hour, location, and route.",
        "A technical pipeline is not enough; results must be easy to demonstrate.",
    ], 1.0, 2.78, 4.9, 2.35, 13, INK, 0.55)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.9, 2.0, 5.72, 4.32, LIGHT_BLUE, LIGHT_BLUE)
    text(slide, "PROJECT OBJECTIVES", 7.2, 2.35, 2.5, 0.2, 9, BLUE, True)
    bullet_list(slide, [
        "Build a reproducible Spark ETL and analytics pipeline.",
        "Store dashboard-ready aggregates in MongoDB.",
        "Expose results through documented REST endpoints.",
        "Present insights through a responsive web dashboard.",
    ], 7.2, 2.78, 4.9, 2.35, 13, INK, 0.55)

    # 4. Dataset
    slide = prs.slides.add_slide(blank)
    base(slide, 4)
    title(slide, "03 · data source", "NYC TLC Yellow Taxi trip records", "A real public Parquet dataset was downloaded and processed locally with Apache Spark.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 2.0, 5.1, 3.8, NAVY, NAVY)
    text(slide, "JANUARY 2024", 1.05, 2.38, 2, 0.2, 10, MINT, True)
    text(slide, "47.6 MB", 1.02, 2.82, 3.5, 0.6, 34, WHITE, True, "Aptos Display")
    text(slide, "Parquet file\n2.87 million valid records\n19 source columns", 1.05, 3.67, 3.8, 1.0, 16, RGBColor(181, 201, 205))
    text(slide, "Source: NYC Taxi & Limousine Commission (TLC)", 1.05, 5.27, 4.15, 0.18, 8, MINT)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.25, 2.0, 6.38, 3.8, WHITE, LINE)
    text(slide, "IMPORTANT FIELDS USED", 6.6, 2.38, 3, 0.2, 9, TEAL, True)
    fields = [("tpep_pickup_datetime", "Time analysis"), ("tpep_dropoff_datetime", "Duration"), ("trip_distance", "Distance"), ("fare_amount", "Fare"), ("total_amount", "Revenue"), ("PULocationID / DOLocationID", "Routes")]
    for i, (field, use) in enumerate(fields):
        yy = 2.82 + i * 0.44
        text(slide, field, 6.6, yy, 3.2, 0.2, 10, INK, True, "Courier New")
        text(slide, use, 10.1, yy, 1.8, 0.2, 10, MUTED)
        line(slide, 6.6, yy + 0.28, 12.18, yy + 0.28, LINE, 0.7)

    # 5. Architecture
    slide = prs.slides.add_slide(blank)
    base(slide, 5)
    title(slide, "04 · system architecture", "One pipeline, five layers", "Each component has a focused responsibility: process, store, serve, visualize, and communicate.")
    labels = [("1", "TLC Parquet", "Raw public data", LIGHT_BLUE, BLUE), ("2", "Apache Spark", "ETL + aggregations", PALE_MINT, TEAL), ("3", "MongoDB", "Analytics store", LIGHT_ORANGE, ORANGE), ("4", "FastAPI", "REST backend", RGBColor(238, 233, 253), PURPLE), ("5", "Dashboard", "Charts + KPIs", NAVY, MINT)]
    xs = [0.65, 3.15, 5.65, 8.15, 10.65]
    for i, (num, name, desc, fill, accent) in enumerate(labels):
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, xs[i], 2.55, 2.02, 1.72, fill, fill)
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, xs[i] + 0.16, 2.75, 0.36, 0.36, accent, accent)
        text(slide, num, xs[i] + 0.16, 2.79, 0.36, 0.18, 9, WHITE if accent != MINT else NAVY, True, "Aptos", PP_ALIGN.CENTER)
        text(slide, name, xs[i] + 0.16, 3.3, 1.7, 0.28, 15, WHITE if fill == NAVY else INK, True)
        text(slide, desc, xs[i] + 0.16, 3.72, 1.7, 0.22, 9, MINT if fill == NAVY else MUTED)
        if i < 4:
            line(slide, xs[i] + 2.04, 3.41, xs[i] + 2.42, 3.41, TEAL, 2, True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.05, 5.15, 11.15, 0.92, WHITE, LINE)
    text(slide, "Design principle", 1.35, 5.43, 1.7, 0.2, 9, TEAL, True)
    text(slide, "Process expensive raw data once → store small, fast-to-query results → serve a responsive user experience.", 3.1, 5.34, 8.55, 0.35, 16, INK, True, "Aptos Display")

    # 6. Spark ETL
    slide = prs.slides.add_slide(blank)
    base(slide, 6)
    title(slide, "05 · distributed processing", "What Apache Spark does", "Spark reads the Parquet file, applies data-quality rules, derives analytical fields, and performs distributed aggregations.")
    add_code(slide, "raw = spark.read.parquet(input)\ntrips = (raw.select(...)\n  .withColumn(\"duration_minutes\", ...)\n  .withColumn(\"pickup_hour\", F.hour(...))\n  .filter(validity_rules)\n  .cache())", 0.7, 2.0, 5.42, 3.8, 11)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.55, 2.0, 6.08, 3.8, WHITE, LINE)
    text(slide, "PROCESSING STEPS", 6.9, 2.36, 2, 0.2, 9, TEAL, True)
    steps = [
        ("01", "Read", "Load Parquet columns with Spark"),
        ("02", "Clean", "Remove invalid distance, fare, zones, and durations"),
        ("03", "Derive", "Calculate duration and pickup hour"),
        ("04", "Aggregate", "Group by hour, route, and pickup zone"),
        ("05", "Load", "Write compact results to MongoDB"),
    ]
    for i, (num, name, desc) in enumerate(steps):
        yy = 2.78 + i * 0.55
        text(slide, num, 6.9, yy, 0.35, 0.2, 9, TEAL, True, "Courier New")
        text(slide, name, 7.45, yy, 1.1, 0.2, 11, INK, True)
        text(slide, desc, 8.55, yy, 3.5, 0.2, 10, MUTED)
    text(slide, "Result: 2,868,178 valid records", 0.92, 6.12, 5.0, 0.25, 15, TEAL, True)
    text(slide, "The pipeline runs locally with master=local[*], using all available CPU cores.", 6.9, 6.12, 5.4, 0.25, 10, MUTED)

    # 7. MongoDB
    slide = prs.slides.add_slide(blank)
    base(slide, 7)
    title(slide, "06 · analytical storage", "MongoDB stores dashboard-ready collections", "The application does not scan millions of raw trip records for each page load; it reads compact aggregates instead.")
    collections = [
        ("overview", "KPI cards", BLUE),
        ("hourly_demand", "24-hour demand", TEAL),
        ("top_routes", "Route ranking", ORANGE),
        ("zone_performance", "Zone metrics", PURPLE),
        ("metadata", "Lineage + status", GREEN),
    ]
    for i, (name, desc, accent) in enumerate(collections):
        yy = 2.02 + i * 0.72
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, yy, 5.15, 0.52, WHITE, LINE)
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.75, yy, 0.08, 0.52, accent)
        text(slide, name, 1.05, yy + 0.13, 2.35, 0.2, 11, INK, True, "Courier New")
        text(slide, desc, 3.7, yy + 0.13, 1.75, 0.2, 10, MUTED)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.55, 2.02, 6.08, 4.12, NAVY, NAVY)
    text(slide, "EXAMPLE DOCUMENT", 6.9, 2.37, 2.2, 0.2, 9, MINT, True)
    add_code(slide, '{\n  "total_trips": 2868178,\n  "total_revenue": 78417975.23,\n  "average_fare": 18.49,\n  "average_distance": 3.29,\n  "average_duration": 14.97\n}', 6.85, 2.83, 5.3, 2.18, 13)
    text(slide, "MongoDB Compass can be used to inspect these collections during the presentation.", 6.9, 5.46, 5.0, 0.42, 11, RGBColor(181, 201, 205))

    # 8. API
    slide = prs.slides.add_slide(blank)
    base(slide, 8)
    title(slide, "07 · backend service", "FastAPI connects analytics to the interface", "The backend provides documented REST endpoints and serves the frontend from one local application.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 2.05, 4.2, 4.1, NAVY, NAVY)
    text(slide, "FASTAPI", 1.05, 2.42, 1.5, 0.2, 10, MINT, True)
    text(slide, "http://127.0.0.1:8000", 1.05, 2.9, 3.4, 0.3, 17, WHITE, True, "Courier New")
    text(slide, "The same service provides:\n\n• Web dashboard at /\n• Interactive API docs at /docs\n• JSON endpoints under /api", 1.05, 3.55, 3.2, 1.55, 14, RGBColor(181, 201, 205))
    endpoints = [
        ("GET", "/api/health", "MongoDB status"),
        ("GET", "/api/overview", "KPI metrics"),
        ("GET", "/api/hourly-demand", "Hourly chart"),
        ("GET", "/api/top-routes", "Route table"),
        ("GET", "/api/zone-performance", "Zone chart"),
    ]
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 5.35, 2.05, 7.28, 4.1, WHITE, LINE)
    text(slide, "REST API CONTRACT", 5.7, 2.42, 2.3, 0.2, 9, TEAL, True)
    for i, (method, endpoint, desc) in enumerate(endpoints):
        yy = 2.86 + i * 0.55
        label_box(slide, method, 5.7, yy - 0.03, 0.62, PALE_MINT, TEAL, 8)
        text(slide, endpoint, 6.55, yy + 0.09, 2.55, 0.2, 11, INK, True, "Courier New")
        text(slide, desc, 9.55, yy + 0.09, 2.2, 0.2, 10, MUTED)
        line(slide, 5.7, yy + 0.42, 12.2, yy + 0.42, LINE, 0.7)

    # 9. Frontend
    slide = prs.slides.add_slide(blank)
    base(slide, 9)
    title(slide, "08 · frontend", "A presentation-ready analytics dashboard", "The browser fetches JSON from FastAPI and turns it into KPI cards, charts, and a route ranking table.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.65, 1.95, 12.02, 4.25, WHITE, LINE)
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.65, 1.95, 12.02, 0.58, NAVY)
    text(slide, "URBAN MOBILITY ANALYTICS", 0.98, 2.14, 4.0, 0.2, 11, WHITE, True)
    label_box(slide, "MongoDB connected", 10.35, 2.05, 1.85, PALE_MINT, TEAL, 8)
    for i, (label, value, accent) in enumerate([("TOTAL TRIPS", "2.87M", BLUE), ("REVENUE", "$78.42M", ORANGE), ("AVG FARE", "$18.49", GREEN), ("AVG DISTANCE", "3.29 mi", PURPLE)]):
        xx = 0.98 + i * 2.9
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, xx, 2.88, 2.55, 0.88, PAPER, LINE)
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, xx, 2.88, 0.05, 0.88, accent)
        text(slide, label, xx + 0.18, 3.06, 2.1, 0.13, 7, MUTED, True)
        text(slide, value, xx + 0.18, 3.29, 2.1, 0.25, 16, INK, True)
    text(slide, "Trips by pickup hour", 1.0, 4.22, 2.4, 0.2, 10, INK, True)
    add_line_chart(slide, [28, 22, 18, 15, 14, 22, 35, 48, 55, 50, 42, 37, 34, 32, 31, 35, 43, 57, 64, 60, 51, 44, 39, 33], ["00", "03", "06", "09", "12", "15", "18", "21", "23"], 0.95, 4.52, 5.35, 1.25, BLUE)
    text(slide, "Busiest zones", 6.7, 4.22, 2.0, 0.2, 10, INK, True)
    add_bar_chart(slide, [94, 88, 76, 70, 65, 62], ["132", "138", "161", "237", "170", "234"], 6.65, 4.48, 5.55, 1.3, GREEN, "trips")

    # 10. Results
    slide = prs.slides.add_slide(blank)
    base(slide, 10)
    title(slide, "09 · results", "What the processed data tells us", "The January 2024 results show clear demand patterns and provide a compact analytical view for decision-making.")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 1.95, 5.7, 4.25, WHITE, LINE)
    text(slide, "HOURLY DEMAND PATTERN", 1.03, 2.3, 2.8, 0.2, 9, TEAL, True)
    text(slide, "Demand rises through the day and is strongest in the evening period.", 1.03, 2.62, 4.75, 0.36, 13, INK, True)
    add_line_chart(slide, [25, 18, 14, 12, 13, 19, 31, 45, 53, 48, 40, 35, 33, 31, 30, 34, 42, 55, 63, 59, 50, 43, 38, 31], ["00", "03", "06", "09", "12", "15", "18", "21", "23"], 1.05, 3.25, 4.9, 2.15, BLUE)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.7, 1.95, 5.92, 4.25, PALE_MINT, PALE_MINT)
    text(slide, "KEY FINDINGS", 7.03, 2.3, 2.0, 0.2, 9, TEAL, True)
    bullet_list(slide, [
        "2.87M trips remained after quality filtering.",
        "$78.42M in total trip revenue was aggregated.",
        "Average trip duration was 14.97 minutes.",
        "Zone 132 was the busiest pickup area in the result set.",
        "Route and hourly aggregates are ready for repeated API queries.",
    ], 7.03, 2.75, 5.0, 2.6, 13, INK, 0.55)

    # 11. Demonstration flow
    slide = prs.slides.add_slide(blank)
    base(slide, 11)
    title(slide, "10 · live demonstration", "How to present the system", "A short sequence shows the complete path from processing command to user-facing insight.")
    demo_steps = [
        ("01", "Show MongoDB", "Open Compass and show the taxi_analytics collections."),
        ("02", "Show API docs", "Open http://127.0.0.1:8000/docs and expand an endpoint."),
        ("03", "Open dashboard", "Show KPI cards, hourly demand, zones, and routes."),
        ("04", "Explain Spark", "Open spark_analyze.py and point to cleaning + groupBy."),
        ("05", "Discuss insight", "Connect the visual pattern to transport planning."),
    ]
    for i, (num, name, desc) in enumerate(demo_steps):
        yy = 2.0 + i * 0.78
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 0.85, yy, 0.48, 0.48, TEAL, TEAL)
        text(slide, num, 0.85, yy + 0.13, 0.48, 0.2, 9, WHITE, True, "Courier New", PP_ALIGN.CENTER)
        text(slide, name, 1.65, yy + 0.04, 2.1, 0.23, 15, INK, True)
        text(slide, desc, 4.05, yy + 0.07, 7.8, 0.25, 12, MUTED)
        if i < len(demo_steps) - 1:
            line(slide, 1.09, yy + 0.5, 1.09, yy + 0.78, LINE, 1.5)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.95, 5.82, 2.4, 0.47, NAVY, NAVY)
    text(slide, "Keep the browser open", 10.08, 5.95, 2.14, 0.18, 9, MINT, True, "Aptos", PP_ALIGN.CENTER)

    # 12. Challenges and solutions
    slide = prs.slides.add_slide(blank)
    base(slide, 12)
    title(slide, "11 · engineering lessons", "Challenges solved during implementation", "The final system is reproducible because the data-quality and platform issues were handled in code and documented.")
    headers = [("Challenge", 0.8, ORANGE), ("Solution", 6.25, TEAL)]
    for label, x, color in headers:
        text(slide, label.upper(), x, 2.05, 2.4, 0.2, 9, color, True)
    rows = [
        ("Timestamp type mismatch in Spark", "Used Spark unix_timestamp() for TIMESTAMP_NTZ values."),
        ("Malformed historical timestamp values", "Added a year-quality filter before computing metrics."),
        ("Large raw file for web requests", "Precomputed aggregates and stored them in MongoDB."),
        ("Need for an immediate presentation demo", "Added seed_demo.py, then replaced demo data with real Spark output."),
        ("Local Windows Spark warnings", "Warnings about winutils/native Hadoop libraries were non-blocking; job completed."),
    ]
    for i, (challenge, solution) in enumerate(rows):
        yy = 2.48 + i * 0.68
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, yy, 5.0, 0.5, WHITE, LINE)
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.2, yy, 6.35, 0.5, PALE_MINT, PALE_MINT)
        text(slide, challenge, 1.0, yy + 0.13, 4.5, 0.2, 10, INK, True)
        text(slide, solution, 6.47, yy + 0.13, 5.8, 0.2, 10, INK)

    # 13. Future work
    slide = prs.slides.add_slide(blank)
    base(slide, 13)
    title(slide, "12 · next steps", "How this project can grow", "The current system is a strong batch analytics foundation for more advanced big data features.")
    future = [
        ("01", "Zone names + maps", "Join taxi zone lookup data and add a geographic heatmap."),
        ("02", "Spark ML", "Predict fare or trip duration from distance, hour, and zone."),
        ("03", "Streaming", "Use Kafka or Spark Structured Streaming for live demand."),
        ("04", "More context", "Add weather, holidays, events, and traffic data."),
    ]
    for i, (num, name, desc) in enumerate(future):
        xx = 0.75 + (i % 2) * 6.05
        yy = 2.08 + (i // 2) * 1.62
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, xx, yy, 5.45, 1.2, WHITE, LINE)
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, xx + 0.25, yy + 0.25, 0.65, 0.65, TEAL, TEAL)
        text(slide, num, xx + 0.25, yy + 0.46, 0.65, 0.17, 9, WHITE, True, "Courier New", PP_ALIGN.CENTER)
        text(slide, name, xx + 1.15, yy + 0.24, 3.9, 0.24, 15, INK, True)
        text(slide, desc, xx + 1.15, yy + 0.61, 3.9, 0.3, 10, MUTED)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.8, 5.7, 9.7, 0.55, NAVY, NAVY)
    text(slide, "Batch analytics today → intelligent urban mobility platform tomorrow", 2.0, 5.87, 9.3, 0.2, 14, MINT, True, "Aptos Display", PP_ALIGN.CENTER)

    # 14. Summary
    slide = prs.slides.add_slide(blank)
    base(slide, 14, True)
    title(slide, "13 · conclusion", "A complete big data analytics application", "The project combines distributed computing, NoSQL storage, backend engineering, and frontend communication.", True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.75, 2.1, 11.85, 2.55, RGBColor(28, 62, 76), RGBColor(46, 91, 105))
    bullet_list(slide, [
        "Apache Spark processed and cleaned 2.87 million taxi trips.",
        "MongoDB stores reusable, dashboard-ready analytical collections.",
        "FastAPI exposes the results through a clean REST interface.",
        "The dashboard makes demand, revenue, zones, and routes easy to explain.",
    ], 1.15, 2.55, 10.7, 1.75, 16, WHITE, 0.42)
    text(slide, "Thank you", 0.75, 5.42, 4.0, 0.55, 34, MINT, True, "Aptos Display")
    text(slide, "Questions?", 0.78, 6.1, 2.0, 0.25, 14, RGBColor(181, 201, 205))
    text(slide, "Jash Parekh · 23BIT033", 9.15, 6.1, 3.5, 0.25, 12, MINT, True, "Aptos", PP_ALIGN.RIGHT)

    # 15. References / run commands
    slide = prs.slides.add_slide(blank)
    base(slide, 15)
    title(slide, "14 · appendix", "Run commands and project references", "Use these commands during a technical demonstration or to reproduce the result on another machine.")
    add_code(slide, "cd big-data-taxi-analytics\n.venv\\Scripts\\Activate.ps1\npython scripts\\download_data.py --year 2024 --month 1\npython scripts\\spark_analyze.py --input data/raw/yellow_tripdata_2024-01.parquet\nuvicorn backend.main:app --reload", 0.75, 2.0, 6.0, 3.45, 11)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 7.1, 2.0, 5.5, 3.45, WHITE, LINE)
    text(slide, "PROJECT FILES", 7.45, 2.38, 2, 0.2, 9, TEAL, True)
    bullet_list(slide, [
        "scripts/spark_analyze.py — Spark ETL",
        "backend/main.py — FastAPI routes",
        "backend/static/ — dashboard frontend",
        "scripts/seed_demo.py — demo data",
        "docs/project_report.md — written report",
    ], 7.45, 2.78, 4.5, 2.05, 11, INK, 0.37)
    text(slide, "Dashboard: http://127.0.0.1:8000", 0.78, 5.95, 5.4, 0.24, 13, TEAL, True, "Courier New")
    text(slide, "API docs: http://127.0.0.1:8000/docs", 7.15, 5.95, 5.0, 0.24, 13, TEAL, True, "Courier New")

    prs.core_properties.title = "Urban Mobility Big Data Analytics Dashboard"
    prs.core_properties.subject = "Big Data Analytics university project presentation"
    prs.core_properties.author = "Jash Parekh"
    prs.core_properties.keywords = "Apache Spark, MongoDB, FastAPI, Big Data, Taxi Analytics"
    prs.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build_presentation()
